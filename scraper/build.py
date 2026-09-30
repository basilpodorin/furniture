"""Сводит raw/*.json в базу актуального велюра: Excel (с миниатюрами), CSV, JSON и HTML-галерею.

В основную базу попадают только цвета в наличии (in_stock = True). Аметист, у которого
наличие проверить нельзя, выводится отдельно с пометкой «наличие не проверено».
"""
import csv
import html
import io
import json
import os
import re
from collections import OrderedDict, defaultdict

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

from common import DATA_DIR, PHOTO_DIR, PHOTO_MAX_SIDE, PHOTO_QUALITY, RAW_DIR

SOURCES = ["souz_m", "artex", "vip_textil", "ametist"]
SUPPLIER_ORDER = ["Союз-М", "Артекс", "Вип Текстиль", "Аметист"]
UNVERIFIED = "Аметист"  # наличие не проверить: официальный сайт закрыт капчей
STOCK_CHECK = {
    "Союз-М": "«Остатки тканей» souz-m.ru/leftovers (Москва, в отрез): «Да»",
    "Артекс": "Статус «В наличии» на artextkani.ru",
    "Вип Текстиль": "«Актуальные остатки» viptextil.ru: «есть в наличии»",
    "Аметист": "Не проверено: ametist-store.ru закрыт капчей",
}

# (ключ, заголовок, ширина колонки)
COLUMNS = [
    ("supplier", "Поставщик", 14),
    ("collection", "Коллекция", 18),
    ("article", "Артикул / цвет", 26),
    ("color", "Цвет (как на сайте)", 18),
    ("status", "Статус", 13),
    ("composition", "Состав", 22),
    ("width_cm", "Ширина, см", 10),
    ("density", "Плотность", 16),
    ("martindale", "Истирание (Мартиндейл), циклов", 14),
    ("martindale_text", "Истирание (как на сайте)", 18),
    ("country", "Страна", 10),
    ("price_rub", "Цена, ₽/п.м.", 11),
    ("old_price_rub", "Цена без скидки, ₽", 11),
    ("price_note", "Комментарий к цене", 26),
    ("stock", "Наличие", 18),
    ("url", "Страница товара", 40),
    ("photo_file", "Фото (файл в папке velour)", 40),
    ("photo_url", "Фото (оригинал на сайте)", 40),
    ("source", "Источник", 30),
    ("collected", "Дата сбора", 11),
]
HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="4B3B6B")
WRAP = Alignment(wrap_text=True, vertical="center")


def load():
    rows = []
    for name in SOURCES:
        path = os.path.join(RAW_DIR, name + ".json")
        if not os.path.exists(path):
            print(f"нет {path}, пропускаю")
            continue
        for r in json.load(open(path, encoding="utf-8")):
            # Артекс: карточки-«обложки» коллекций без цвета и фото — не позиции
            if name == "artex" and not r.get("collection"):
                continue
            if name == "souz_m" and r.get("density_gsm") is not None:
                r["density"] = f"{str(r['density_gsm']).replace('.', ',')} г/м²"
            if r.get("martindale") and not r.get("martindale_text"):
                r["martindale_text"] = f"{r['martindale']:,} циклов".replace(",", " ")
            rows.append(r)
    rows.sort(key=lambda r: (SUPPLIER_ORDER.index(r["supplier"]),
                             (r["collection"] or "").lower(), _natural(r["article"])))
    return rows


def _natural(s):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", s or "")]


def normalize_photos(rows):
    """Приводит фото к длинной стороне <= PHOTO_MAX_SIDE."""
    n = 0
    for r in rows:
        p = os.path.join(DATA_DIR, r["photo_file"])
        im = Image.open(p)
        if max(im.size) > PHOTO_MAX_SIDE:
            im = im.convert("RGB")
            im.thumbnail((PHOTO_MAX_SIDE, PHOTO_MAX_SIDE), Image.LANCZOS)
            im.save(p, "JPEG", quality=PHOTO_QUALITY, optimize=True, progressive=True)
            n += 1
    if n:
        print(f"фото уменьшено до {PHOTO_MAX_SIDE}px: {n}")


def prune_photos(rows):
    """Удаляет из photos/ файлы, которых нет в базе (например, цвета не в наличии)."""
    keep = {os.path.normpath(os.path.join(DATA_DIR, r["photo_file"])) for r in rows}
    removed = 0
    for root, _, files in os.walk(PHOTO_DIR, topdown=False):
        for f in files:
            p = os.path.normpath(os.path.join(root, f))
            if p not in keep:
                os.remove(p)
                removed += 1
        if not os.listdir(root):
            os.rmdir(root)
    print(f"удалено фото вне базы: {removed}")


def write_csv_json(rows, name):
    keys = [k for k, _, _ in COLUMNS]
    with open(os.path.join(DATA_DIR, name + ".csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([h for _, h, _ in COLUMNS])
        for r in rows:
            w.writerow(["" if r.get(k) is None else r.get(k) for k in keys])
    with open(os.path.join(DATA_DIR, name + ".json"), "w", encoding="utf-8") as f:
        json.dump([{k: r.get(k) for k in keys} for r in rows], f, ensure_ascii=False, indent=1)


def collections(rows):
    groups = OrderedDict()
    for r in rows:
        groups.setdefault((r["supplier"], r["collection"]), []).append(r)
    out = []
    for (sup, col), items in groups.items():
        def uniq(k):
            vals = sorted({str(i[k]) for i in items if i.get(k) not in (None, "")})
            return "; ".join(vals) if vals else None
        prices = [i["price_rub"] for i in items if i.get("price_rub")]
        out.append({
            "supplier": sup, "collection": col, "count": len(items),
            "composition": uniq("composition"), "width_cm": uniq("width_cm"),
            "density": uniq("density"), "martindale_text": uniq("martindale_text"),
            "country": uniq("country"), "status": uniq("status"),
            "price": (f"{min(prices)}" if min(prices) == max(prices) else f"{min(prices)}–{max(prices)}") if prices else uniq("price_note"),
            "url": items[0]["url"],
        })
    return out


def thumb_bytes(path, size=64):
    im = Image.open(path).convert("RGB")
    im.thumbnail((size, size), Image.LANCZOS)
    b = io.BytesIO()
    im.save(b, "JPEG", quality=80)
    b.seek(0)
    return b


def _header(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD_FONT, HEAD_FILL, WRAP
    ws.row_dimensions[1].height = 42


def add_items_sheet(wb, title, rows):
    """Лист «одна строка — один цвет» с миниатюрой фото в колонке A."""
    ws = wb.create_sheet(title)
    _header(ws, ["Фото"] + [h for _, h, _ in COLUMNS], [10] + [w for _, _, w in COLUMNS])
    link_cols = {k: i for i, (k, _, _) in enumerate(COLUMNS, start=2) if k in ("url", "photo_url", "photo_file")}
    for n, r in enumerate(rows, start=2):
        ws.append([None] + [r.get(k) for k, _, _ in COLUMNS])
        ws.row_dimensions[n].height = 50
        for ci in range(2, len(COLUMNS) + 2):
            ws.cell(row=n, column=ci).alignment = WRAP
        for k, ci in link_cols.items():
            if r.get(k):
                cell = ws.cell(row=n, column=ci)
                cell.hyperlink = r[k]
                cell.style = "Hyperlink"
        ws.add_image(XLImage(thumb_bytes(os.path.join(DATA_DIR, r["photo_file"]))), f"A{n}")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS) + 1)}{len(rows) + 1}"


def write_xlsx(current, unverified, all_rows):
    wb = Workbook()

    ws = wb.active
    ws.title = "Сводка"
    _header(ws, ["Поставщик", "Коллекций в наличии", "Цветов в наличии", "Всего цветов на сайте",
                 "Как проверено наличие", "Источник данных"], [16, 13, 12, 13, 50, 60])
    total = defaultdict(int)
    for r in all_rows:
        total[r["supplier"]] += 1
    by = OrderedDict((s, [r for r in current if r["supplier"] == s]) for s in SUPPLIER_ORDER if s != UNVERIFIED)
    for sup, items in by.items():
        ws.append([sup, len({i["collection"] for i in items}), len(items), total[sup],
                   STOCK_CHECK[sup], items[0]["source"] if items else None])
    ws.append(["Итого в наличии", len({(r["supplier"], r["collection"]) for r in current}), len(current),
               sum(total[s] for s in by), None, None])
    ws.append([])
    ws.append([f"{UNVERIFIED} (отдельный лист)", len({r["collection"] for r in unverified}), None,
               len(unverified), STOCK_CHECK[UNVERIFIED], unverified[0]["source"] if unverified else None])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP

    add_items_sheet(wb, "Велюр в наличии", current)

    wc = wb.create_sheet("Коллекции")
    cols_def = [("supplier", "Поставщик", 14), ("collection", "Коллекция", 22), ("count", "Цветов в наличии", 10),
                ("status", "Статусы", 16), ("composition", "Состав", 26), ("width_cm", "Ширина, см", 10),
                ("density", "Плотность", 18), ("martindale_text", "Истирание", 20), ("country", "Страна", 10),
                ("price", "Цена, ₽/п.м.", 16), ("url", "Ссылка", 50)]
    _header(wc, [h for _, h, _ in cols_def], [w for _, _, w in cols_def])
    cols = collections(current)
    for n, c in enumerate(cols, start=2):
        wc.append([c.get(k) for k, _, _ in cols_def])
        wc.cell(row=n, column=len(cols_def)).hyperlink = c["url"]
        wc.cell(row=n, column=len(cols_def)).style = "Hyperlink"
    wc.freeze_panes = "C2"
    wc.auto_filter.ref = f"A1:{get_column_letter(len(cols_def))}{len(cols) + 1}"

    if unverified:
        add_items_sheet(wb, f"{UNVERIFIED} (не проверено)", unverified)

    wb.active = 0
    wb.save(os.path.join(DATA_DIR, "velour.xlsx"))
    return cols


def write_gallery(current, unverified):
    e = html.escape

    def section(rows):
        by = OrderedDict()
        for r in rows:
            by.setdefault(r["collection"], []).append(r)
        out = []
        for col, items in by.items():
            i0 = items[0]
            spec = ", ".join(filter(None, [i0.get("composition"), i0.get("density"), i0.get("martindale_text"),
                                           f"{i0['width_cm']} см" if i0.get("width_cm") else None]))
            out.append(f"<h3>{e(col or '')} <small>{len(items)} цв. · {e(spec)}</small></h3><div class=g>")
            for r in items:
                src = e(r["photo_file"])
                extra = [r["color"]] if r.get("color") and r["color"] not in (r["article"] or "") else []
                if r.get("status"):
                    extra.append(r["status"])
                price = f"{r['price_rub']} ₽" if r.get("price_rub") else ""
                out.append(f'<figure><a href="{src}" target="_blank"><img loading="lazy" src="{src}" alt=""></a>'
                           f'<figcaption><a href="{e(r["url"])}" target="_blank">{e(r["article"] or "")}</a>'
                           f'{"<br>" + e(" · ".join(extra)) if extra else ""}'
                           f'{"<br>" + price if price else ""}</figcaption></figure>')
            out.append("</div>")
        return out

    groups = OrderedDict((s, [r for r in current if r["supplier"] == s]) for s in SUPPLIER_ORDER if s != UNVERIFIED)
    parts = ["""<!doctype html><html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Велюр в наличии</title>
<style>
body{font-family:system-ui,sans-serif;margin:0 16px 40px;background:#faf9f7;color:#222}
h1{margin:20px 0 4px}h2{margin:32px 0 8px;border-bottom:2px solid #4b3b6b;padding-bottom:4px}
h3{margin:18px 0 6px;font-size:16px}h3 small{color:#777;font-weight:normal}
nav a{margin-right:14px}.note{background:#fff4d6;border:1px solid #e8c96a;padding:8px 12px;border-radius:6px}
.g{display:grid;grid-template-columns:repeat(auto-fill,minmax(140px,1fr));gap:10px}
figure{margin:0;background:#fff;border:1px solid #e3e0da;border-radius:6px;overflow:hidden}
figure img{width:100%;aspect-ratio:1;object-fit:cover;display:block;background:#eee}
figcaption{font-size:12px;padding:4px 6px;line-height:1.3}
figcaption a{color:inherit}
</style></head><body>""",
             f"<h1>Велюр в наличии</h1><p>{len(current)} цветов · данные на "
             f"{e(current[0]['collected']) if current else ''}</p><nav>"]
    for sup, rows in groups.items():
        parts.append(f'<a href="#{e(sup)}">{e(sup)} ({len(rows)})</a>')
    if unverified:
        parts.append(f'<a href="#{e(UNVERIFIED)}">{e(UNVERIFIED)} — не проверено ({len(unverified)})</a>')
    parts.append("</nav>")
    for sup, rows in groups.items():
        parts.append(f'<h2 id="{e(sup)}">{e(sup)}</h2>')
        parts += section(rows)
    if unverified:
        parts.append(f'<h2 id="{e(UNVERIFIED)}">{e(UNVERIFIED)} — наличие не проверено</h2>'
                     f'<p class=note>Сайт Аметиста закрыт капчей, данные взяты у дилера в Беларуси. '
                     f'Наличие уточняйте у Аметиста.</p>')
        parts += section(unverified)
    parts.append("</body></html>")
    open(os.path.join(DATA_DIR, "gallery.html"), "w", encoding="utf-8").write("\n".join(parts))


def main():
    rows = load()
    current = [r for r in rows if r.get("in_stock") and r["supplier"] != UNVERIFIED]
    unverified = [r for r in rows if r["supplier"] == UNVERIFIED]
    missing = [r["article"] for r in current + unverified if not r.get("photo_file")]
    if missing:
        raise SystemExit(f"нет фото у {len(missing)} позиций: {missing[:10]}")
    normalize_photos(current + unverified)
    prune_photos(current + unverified)
    write_csv_json(current, "velour")
    write_csv_json(unverified, "ametist_unverified")
    cols = write_xlsx(current, unverified, rows)
    write_gallery(current, unverified)
    by = defaultdict(int)
    for r in current:
        by[r["supplier"]] += 1
    print("в наличии:", dict(by), "всего", len(current), "коллекций", len(cols),
          f"| {UNVERIFIED} (не проверено): {len(unverified)}")


if __name__ == "__main__":
    main()
